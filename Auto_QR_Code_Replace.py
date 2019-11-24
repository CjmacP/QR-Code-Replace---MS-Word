import qrcode
import docx
from docx.shared import Inches

def qr_code_maker(url):

    qr = qrcode.make(url)
    qr.save(#url name here# + '.png')
    
def replace_QR(new_url_location):
    
    runs = doc2.tables[1].rows[0].cells[0].paragraphs[0].runs[0]

    doc2.tables[1].rows[0].cells[0]._element.clear_content()

    p = doc2.tables[1].rows[0].cells[0].add_paragraph()
    r = p.add_run()
    loc = new_url_location  # Location from line 7

    r.add_picture(loc, width=Inches(1.6), height=Inches(1.6))

    doc2.tables[i].rows[0].cells[0].paragraphs[0].runs[0] = runs
    
    
